import json
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from src.renderers.docx import DocxRenderer

FIXTURES_DIR = Path(__file__).parent / "fixtures"


def _load_template() -> bytes:
    return (FIXTURES_DIR / "test_template.docx").read_bytes()


def _load_data() -> dict:
    return json.loads((FIXTURES_DIR / "test_data.json").read_text())


def _extract_text(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


class TestDocxRendererHappyPath:
    def test_simple_placeholders_replaced(self) -> None:
        renderer = DocxRenderer()

        result = renderer.render(_load_template(), _load_data())

        text = _extract_text(result.content)
        assert "Dupont" in text
        assert "Nice" in text

    def test_conditional_block_rendered(self) -> None:
        renderer = DocxRenderer()

        result = renderer.render(_load_template(), _load_data())

        text = _extract_text(result.content)
        assert "Vous etes administrateur" in text

    def test_iterative_block_rendered(self) -> None:
        renderer = DocxRenderer()

        result = renderer.render(_load_template(), _load_data())

        text = _extract_text(result.content)
        assert "dossier-001" in text
        assert "dossier-002" in text
        assert "dossier-003" in text

    def test_no_warnings_when_all_keys_present(self) -> None:
        renderer = DocxRenderer()

        result = renderer.render(_load_template(), _load_data())

        assert result.warnings == []

    def test_output_is_valid_docx(self) -> None:
        renderer = DocxRenderer()

        result = renderer.render(_load_template(), _load_data())

        with zipfile.ZipFile(BytesIO(result.content), "r") as z:
            names = z.namelist()
            assert "word/document.xml" in names
            assert "[Content_Types].xml" in names


class TestDocxRendererMissingKeys:
    def test_missing_simple_key_keeps_placeholder(self) -> None:
        """DebugUndefined renders `{{ key }}` for missing placeholders."""
        renderer = DocxRenderer()

        result = renderer.render(
            _load_template(),
            {"nom": "Dupont", "is_admin": True, "items": []},
        )

        text = _extract_text(result.content)
        assert "Dupont" in text
        assert "{{ ville }}" in text

    def test_missing_key_generates_warning(self) -> None:
        renderer = DocxRenderer()

        result = renderer.render(
            _load_template(),
            {"nom": "Dupont", "is_admin": True, "items": []},
        )

        assert "La donnée pour 'ville' n'est pas définie" in result.warnings

    def test_multiple_missing_keys(self) -> None:
        renderer = DocxRenderer()

        result = renderer.render(_load_template(), {})

        assert any("nom" in w for w in result.warnings)
        assert any("ville" in w for w in result.warnings)


class TestDocxRendererEdgeCases:
    def test_extra_keys_ignored(self) -> None:
        renderer = DocxRenderer()
        data = {**_load_data(), "unknown_key": "ignored"}

        result = renderer.render(_load_template(), data)

        assert result.warnings == []
        assert "ignored" not in _extract_text(result.content)

    def test_conditional_false_bool_hides_block(self) -> None:
        """Python `False` is falsy in Jinja — no stringification needed."""
        renderer = DocxRenderer()
        data = {**_load_data(), "is_admin": False}

        result = renderer.render(_load_template(), data)

        text = _extract_text(result.content)
        assert "Vous etes administrateur" not in text

    def test_empty_list_produces_no_items(self) -> None:
        renderer = DocxRenderer()
        data = {**_load_data(), "items": []}

        result = renderer.render(_load_template(), data)

        text = _extract_text(result.content)
        assert "dossier" not in text